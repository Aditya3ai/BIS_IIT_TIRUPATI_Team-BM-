#!/usr/bin/env python3
"""
Generate an 8-slide PPTX for the hackathon presentation using python-pptx.
Creates `presentation/bis_demo_presentation.pptx`.
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def add_text_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for b in bullets:
        p = body.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(18)


def main():
    prs = Presentation()
    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "AI-powered BIS Recommendation Engine"
    slide.placeholders[1].text = "RAG-based retrieval + BM25 + FAISS hybrid. Demo & evaluation"

    add_text_slide(prs, "Problem", [
        "Finding relevant BIS/IS standards from product descriptions is slow",
        "Manual mapping is error-prone and inconsistent",
    ])

    add_text_slide(prs, "Solution", [
        "RAG pipeline: ingest PDFs → chunk standards → BM25 + FAISS retrieval",
        "Top-3 recommendations with short LLM rationale (hallucination-guarded)",
    ])

    add_text_slide(prs, "Architecture", [
        "Ingestion: pdfplumber/pypdf → standard-level chunks",
        "Retrieval: BM25 (rank_bm25) + FAISS vector search → RRF fusion",
        "Rationale: Google Gemini / Ollama / keyword fallback",
    ])

    add_text_slide(prs, "Chunking", [
        "1,205 standard-level chunks (data/chunks.json)",
        "Chunk metadata includes `standard_id`, title, page refs",
    ])

    add_text_slide(prs, "Demo", [
        "Run: python inference.py --input sample.json --output out.json",
        "Shows top-3 standards + brief rationale",
    ])

    add_text_slide(prs, "Evaluation", [
        "HitRate@3: 90% (public test)",
        "MRR@5: 0.742; Avg latency <5s",
    ])

    add_text_slide(prs, "Team & Next Steps", [
        "Team: (Your Names)",
        "Next: tighten hallucination guard, polish UI, prepare deployment",
    ])

    out_dir = Path(__file__).resolve().parents[1] / "presentation"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "bis_demo_presentation.pptx"
    prs.save(str(out_path))
    print(f"Wrote presentation to {out_path}")


if __name__ == '__main__':
    main()
